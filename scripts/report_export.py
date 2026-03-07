# scripts/report_export.py
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from datetime import datetime

GREEN  = PatternFill("solid", fgColor="C6EFCE")
RED    = PatternFill("solid", fgColor="FFC7CE")
YELLOW = PatternFill("solid", fgColor="FFEB9C")
BLUE   = PatternFill("solid", fgColor="BDD7EE")
HEADER = PatternFill("solid", fgColor="1F3864")
BOLD   = Font(bold=True)
WHITE  = Font(bold=True, color="FFFFFF")

def style_header_row(ws, row, cols):
    for col in range(1, cols+1):
        cell = ws.cell(row=row, column=col)
        cell.fill   = HEADER
        cell.font   = WHITE
        cell.alignment = Alignment(horizontal="center")

def export_to_excel(profile, check_results, scores, output_path):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Executive Summary ──────────────────────
    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.column_dimensions["A"].width = 25
    ws1.column_dimensions["B"].width = 20

    ws1["A1"] = "DATA QUALITY REPORT — RETAIL PRODUCTS"
    ws1["A1"].font = Font(bold=True, size=14, color="1F3864")
    ws1["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws1["A2"].font = Font(italic=True, color="808080")

    ws1["A4"] = "Dimension"
    ws1["B4"] = "Score (%)"
    style_header_row(ws1, 4, 2)

    rows = [
        ("Completeness",  scores["completeness"]),
        ("Validity",      scores["validity"]),
        ("Uniqueness",    scores["uniqueness"]),
        ("Consistency",   scores["consistency"]),
        ("OVERALL SCORE", scores["overall_score"]),
    ]
    for i, (dim, val) in enumerate(rows, start=5):
        ws1[f"A{i}"] = dim
        ws1[f"B{i}"] = val
        color = GREEN if val >= 75 else (YELLOW if val >= 60 else RED)
        ws1[f"B{i}"].fill = color
        if dim == "OVERALL SCORE":
            ws1[f"A{i}"].font = BOLD
            ws1[f"B{i}"].font = BOLD

    ws1[f"A{len(rows)+6}"] = "Grade"
    ws1[f"B{len(rows)+6}"] = scores["grade"]

    # ── Sheet 2: Data Profile ────────────────────────────
    ws2 = wb.create_sheet("Data Profile")
    headers = ["Column","Data Type","Null Count","Null %",
               "Unique Values","Duplicates","Min","Max","Mean","Outliers"]
    ws2.append(headers)
    style_header_row(ws2, 1, len(headers))

    for col, s in profile.items():
        ws2.append([
            col,
            s.get("data_type",""),
            s.get("null_count",""),
            s.get("null_pct",""),
            s.get("unique_values",""),
            s.get("duplicate_count",""),
            s.get("min","N/A"),
            s.get("max","N/A"),
            s.get("mean","N/A"),
            s.get("outlier_count","N/A")
        ])
        # Highlight high null %
        row = ws2.max_row
        if isinstance(s.get("null_pct"), float) and s["null_pct"] > 10:
            ws2[f"D{row}"].fill = RED

    for col in ws2.columns:
        ws2.column_dimensions[col[0].column_letter].width = 16

    # ── Sheet 3: Quality Checks ──────────────────────────
    ws3 = wb.create_sheet("Quality Checks")
    ws3.append(["Check Key","Column","Rule","Status","Failed Records"])
    style_header_row(ws3, 1, 5)

    for key, r in check_results.items():
        ws3.append([key, r["column"], r["rule"], r["status"], r["failed_count"]])
        row = ws3.max_row
        ws3[f"D{row}"].fill = GREEN if r["passed"] else RED

    for col in ws3.columns:
        ws3.column_dimensions[col[0].column_letter].width = 22

    wb.save(output_path)
    print(f"✅ Excel report saved: {output_path}")

if __name__ == "__main__":
    from ingestion      import load_retail_data
    from profiling      import profile_dataset
    from quality_checks import run_all_checks
    from scoring        import calculate_dq_score
    df      = load_retail_data()
    profile = profile_dataset(df)
    checks  = run_all_checks(df)
    scores  = calculate_dq_score(df, profile, checks)
    export_to_excel(profile, checks, scores, "reports/dq_report.xlsx")